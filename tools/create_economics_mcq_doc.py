import docx
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

def create_economics_docx():
    doc = Document()

    # Set standard margins (0.8 inch)
    for section in doc.sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)

    # Title
    title_p = doc.add_paragraph()
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_run = title_p.add_run('11th & 12th Economics (பொருளியல்)\nமாதிரி வினா-விடைத் தொகுப்பு (MCQ Bank)')
    title_run.font.name = 'Calibri'
    title_run.font.size = Pt(16)
    title_run.font.bold = True
    title_run.font.color.rgb = RGBColor(26, 82, 118)

    # Subtitle / Note
    sub_p = doc.add_paragraph()
    sub_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub_run = sub_p.add_run('பிழை திருத்தப்பட்டு சீரமைக்கப்பட்ட மாதிரி வினாக்கள் (Cleaned & Formatted by JARVIS)')
    sub_run.font.name = 'Calibri'
    sub_run.font.size = Pt(10)
    sub_run.font.italic = True
    sub_run.font.color.rgb = RGBColor(100, 100, 100)

    doc.add_paragraph().paragraph_format.space_after = Pt(10)

    questions = [
        {
            'q_no': '1',
            'ref': '11th Economics - வேலைவாய்ப்புக் கோட்பாடுகள் (Frictional Unemployment)',
            'question': 'உராய்வு வேலையின்மை / பிரிக்சனல் வேலையின்மை (Frictional Unemployment) என்பது கீழ்க்கண்டவற்றுள் எதன் வெளிப்பாடு?',
            'statements': [
                '1. தொழிலாளர்கள் ஓரிடத்திலிருந்து மற்றோர் இடத்திற்கு இடம்பெயர விருப்பம் இல்லாத பொழுது.',
                '2. தகுதிக்கேற்ற வேலைவாய்ப்பு பற்றிய தகவல் பற்றாக்குறை அல்லது அறியாமை நிலவும் பொழுது.'
            ],
            'options': [
                'A) 1 மட்டும் சரி',
                'B) 2 மட்டும் சரி',
                'C) 1 மற்றும் 2 ஆகிய இரண்டும் சரி',
                'D) இரண்டும் தவறு'
            ],
            'answer': 'C) 1 மற்றும் 2 ஆகிய இரண்டும் சரி',
            'explanation': 'உராய்வு வேலையின்மை (Frictional / Search Unemployment) என்பது தொழிலாளர்கள் ஓரிடத்திலிருந்து மற்றொரு இடத்திற்கு இடம்பெயர விருப்பமின்மை, தகுதிக்கேற்ற வேலை குறித்த அறியாமை, மற்றும் தகவல் குறைபாடுகளால் தற்காலிகமாக ஏற்படும் வேலையின்மை ஆகும்.'
        },
        {
            'q_no': '2',
            'ref': '11th Economics - ஊதியம் மற்றும் வேலைவாய்ப்பு வகைகள்',
            'question': 'சரியான கூற்றைத் தேர்ந்தெடு:',
            'statements': [],
            'options': [
                'A) தொழிலாளர்கள் தற்போதைய நடப்புக் கூலியை ஏற்றுக் கொள்ளாமல் தாமாகவே வேலையை விட்டு விலகுவது - தன்னார்வ வேலையின்மை (Voluntary Unemployment).',
                'B) ஒரு பொருளாதாரத்தினுள் உழைப்பாளர்கள் ஒரு வேலையிலிருந்து மற்றொரு வேலைக்கு மாறுவதால் ஏற்படும் தற்காலிக நிலை - உராய்வு வேலையின்மை.',
                'C) நடப்பு கூலி விகிதத்தில் வேலை செய்யத் தயாராக இருந்தும் வேலை கிடைக்காத நிலை - கட்டாய அல்லது விருப்பமில்லா வேலையின்மை (Involuntary Unemployment).',
                'D) மேற்கண்ட கூற்றுகள் அனைத்தும் சரியானவை'
            ],
            'answer': 'D) மேற்கண்ட கூற்றுகள் அனைத்தும் சரியானவை',
            'explanation': 'தொழிலாளர்கள் நிலவும் கூலியில் வேலை செய்ய மறுத்தால் தன்னார்வ வேலையின்மை எனவும், நிலவும் கூலியை ஏற்றுக்கொள்ளத் தயாராக இருந்தும் வேலை கிடைக்காவிடில் அது விருப்பமில்லா வேலையின்மை எனவும் வரையறுக்கப்படுகிறது.'
        },
        {
            'q_no': '3',
            'ref': '12th Economics - பாடம் 5 (பணவியல் பொருளியல் / Monetary Economics)',
            'question': 'சரியான கூற்றைத் தேர்ந்தெடு (வாங்கும் திறன் மற்றும் விலை மட்டம்):',
            'statements': [
                '1. வாங்கும் திறன் (Purchasing Power) என்பது ஒரு அலகு பணம் மூலம் வாங்கக்கூடிய பொருட்கள் அல்லது சேவைகளின் அளவின் அடிப்படையில் வெளிப்படுத்தப்படும் நாணயத்தின் மதிப்பாகும்.',
                '2. விலை மட்டம் உயரும் போது பணத்தின் வாங்கும் திறன் குறைகிறது; விலை மட்டம் குறையும் போது வாங்கும் திறன் அதிகரிக்கிறது.'
            ],
            'options': [
                'A) கூற்று 1 மட்டும் சரி',
                'B) கூற்று 2 மட்டும் சரி',
                'C) கூற்று 1 மற்றும் 2 ஆகிய இரண்டும் சரி',
                'D) இரண்டும் தவறு'
            ],
            'answer': 'C) கூற்று 1 மற்றும் 2 ஆகிய இரண்டும் சரி',
            'explanation': 'பணத்தின் வாங்கும் திறனும் விலை மட்டமும் ஒன்றுக்கொன்று எதிர்மறைத் தொடர்பு (Inverse Relationship) உடையவை. விலை மட்டம் (P) உயரும்போது பணத்தின் வாங்கும் திறன் (1/P) குறையும்.'
        },
        {
            'q_no': '4',
            'ref': '12th Economics - பாடம் 2 (தேசிய வருவாய் / National Income - GDP)',
            'question': 'இந்தியாவில் மொத்த உள்நாட்டு உற்பத்தி (GDP) தொடர்பான கூற்றுகளை ஆராய்க:',
            'statements': [
                '1. ஒரு குறிப்பிட்ட ஆண்டில் நாட்டின் புவியியல் எல்லைக்குள் உற்பத்தி செய்யப்பட்ட அனைத்து இறுதிப் பொருட்கள் மற்றும் பணிகளின் மொத்த பணமதிப்பு GDP எனப்படும்.',
                '2. மொத்த உள்நாட்டு உற்பத்தியில் வெளிநாட்டிலிருந்து ஈட்டப்பட்ட நிகர காரணி வருமானம் (NFIA) சேர்க்கப்படுவதில்லை.'
            ],
            'options': [
                'A) 1 மட்டும் சரி',
                'B) 2 மட்டும் சரி',
                'C) 1 மற்றும் 2 ஆகிய இரண்டும் சரி',
                'D) இரண்டும் தவறு'
            ],
            'answer': 'C) 1 மற்றும் 2 ஆகிய இரண்டும் சரி',
            'explanation': 'GDP என்பது உள்நாட்டு உற்பத்தி மட்டுமே. இதில் வெளிநாட்டு நிகர வருமானம் (NFIA) சேர்க்கப்படும்போது அது மொத்த நாட்டு உற்பத்தி (GNP = GDP + NFIA) ஆக மாறுகிறது.'
        }
    ]

    for q in questions:
        p_q = doc.add_paragraph()
        p_q.paragraph_format.space_before = Pt(10)
        p_q.paragraph_format.space_after = Pt(2)
        
        # Tag
        ref_text = '[' + q['ref'] + ']\n'
        run_tag = p_q.add_run(ref_text)
        run_tag.font.name = 'Calibri'
        run_tag.font.size = Pt(9.5)
        run_tag.font.bold = True
        run_tag.font.color.rgb = RGBColor(180, 80, 0)
        
        # Question
        run_qno = p_q.add_run('வினா ' + q['q_no'] + ': ')
        run_qno.font.bold = True
        run_qno.font.size = Pt(11)
        
        run_q = p_q.add_run(q['question'])
        run_q.font.bold = True
        run_q.font.size = Pt(11)
        
        # Statements if any
        for stmt in q['statements']:
            p_stmt = doc.add_paragraph()
            p_stmt.paragraph_format.left_indent = Inches(0.25)
            p_stmt.paragraph_format.space_after = Pt(2)
            r_stmt = p_stmt.add_run(stmt)
            r_stmt.font.size = Pt(10.5)
            
        # Options
        for opt in q['options']:
            p_opt = doc.add_paragraph()
            p_opt.paragraph_format.left_indent = Inches(0.3)
            p_opt.paragraph_format.space_after = Pt(1)
            r_opt = p_opt.add_run(opt)
            r_opt.font.size = Pt(10.5)
            
        # Answer box / paragraph
        p_ans = doc.add_paragraph()
        p_ans.paragraph_format.left_indent = Inches(0.25)
        p_ans.paragraph_format.space_before = Pt(4)
        p_ans.paragraph_format.space_after = Pt(6)
        
        r_ans_label = p_ans.add_run('விடை: ')
        r_ans_label.font.bold = True
        r_ans_label.font.color.rgb = RGBColor(0, 128, 0)
        
        r_ans = p_ans.add_run(q['answer'] + '\n')
        r_ans.font.bold = True
        r_ans.font.color.rgb = RGBColor(0, 128, 0)
        
        r_exp_label = p_ans.add_run('விளக்கம்: ')
        r_exp_label.font.italic = True
        r_exp_label.font.size = Pt(9.5)
        r_exp_label.font.color.rgb = RGBColor(90, 90, 90)
        
        r_exp = p_ans.add_run(q['explanation'])
        r_exp.font.italic = True
        r_exp.font.size = Pt(9.5)
        r_exp.font.color.rgb = RGBColor(90, 90, 90)
        
        # Divider line
        p_div = doc.add_paragraph()
        p_div.paragraph_format.space_after = Pt(6)
        r_div = p_div.add_run('─' * 55)
        r_div.font.color.rgb = RGBColor(210, 210, 210)

    output_path = r'C:\Users\mukil\Desktop\Economics_MCQ_Questions_Cleaned.docx'
    doc.save(output_path)
    print(f'Successfully saved to {output_path}')

if __name__ == '__main__':
    create_economics_docx()
