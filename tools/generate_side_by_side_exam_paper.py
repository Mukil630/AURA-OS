# -*- coding: utf-8 -*-
"""
Generate professional Tamil Economics Mock Test Question Paper with Side-by-Side (2x2) Options:
a)                b)
c)                d)

- Clean layout for competitive exam practice
- No answers, no explanations, no textbook tags
- Borderless 2x2 table for options
- Converts to PDF via Word COM
- Sends to Mukil's Telegram bot directly in pure professional Tamil
- NO Google Drive upload
"""

import os
import sys
import requests
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
import win32com.client

if sys.stdout.encoding != 'utf-8':
    sys.stdout.reconfigure(encoding='utf-8', errors='replace')

def set_table_no_borders(table):
    """Removes all borders from a docx table."""
    tblPr = table._tbl.tblPr
    tblBorders = OxmlElement('w:tblBorders')
    for b in ['top', 'left', 'bottom', 'right', 'insideH', 'insideV']:
        node = OxmlElement(f'w:{b}')
        node.set(qn('w:val'), 'none')
        tblBorders.append(node)
    tblPr.append(tblBorders)

def set_cell_margins(cell, top=30, bottom=30, left=60, right=60):
    """Sets tight internal cell padding."""
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def extract_questions():
    with open(r"C:\Users\mukil\Desktop\build_test5_doc.py", "r", encoding="utf-8") as f:
        code = f.read()
    
    start_idx = code.find("questions = [")
    end_idx = code.find("    # Filter questions", start_idx)
    if end_idx == -1:
        end_idx = code.find("    # --- Add Questions to Document ---", start_idx)
    
    q_code = code[start_idx:end_idx].strip()
    local_env = {}
    exec(q_code, {}, local_env)
    return local_env["questions"]

def build_exam_paper(docx_path, pdf_path):
    questions = extract_questions()
    print(f"Loaded {len(questions)} questions.")

    doc = docx.Document()

    # Set Margins (0.7 in all sides for optimal 2-column options width)
    for section in doc.sections:
        section.top_margin = Inches(0.7)
        section.bottom_margin = Inches(0.7)
        section.left_margin = Inches(0.7)
        section.right_margin = Inches(0.7)
        section.page_width = Inches(8.27)   # A4
        section.page_height = Inches(11.69)

    # Style
    style = doc.styles['Normal']
    font = style.font
    font.name = 'Nirmala UI'
    font.size = Pt(10.5)
    font.color.rgb = RGBColor(0x11, 0x18, 0x27)

    # --- Header ---
    title_p = doc.add_paragraph()
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_p.paragraph_format.space_before = Pt(0)
    title_p.paragraph_format.space_after = Pt(2)
    t_run = title_p.add_run("பொருளாதாரம் மாதிரித் தேர்வு – 5")
    t_run.font.name = 'Nirmala UI'
    t_run.font.size = Pt(16)
    t_run.font.bold = True
    t_run.font.color.rgb = RGBColor(0x1A, 0x36, 0x5D)

    sub_p = doc.add_paragraph()
    sub_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub_p.paragraph_format.space_after = Pt(6)
    s_run = sub_p.add_run("நேரம்: 1 மணி 30 நிமிடம்          |          மொத்த வினாக்கள்: 75          |          மதிப்பெண்கள்: 75")
    s_run.font.name = 'Nirmala UI'
    s_run.font.size = Pt(10)
    s_run.font.bold = True
    s_run.font.color.rgb = RGBColor(0x4A, 0x55, 0x68)

    div_p = doc.add_paragraph()
    div_p.paragraph_format.space_after = Pt(10)
    div_run = div_p.add_run("―" * 60)
    div_run.font.color.rgb = RGBColor(0xCB, 0xD5, 0xE0)

    # Printable width = 8.27 - 1.4 = 6.87 inches
    col_width = Inches(3.4)

    for item in questions:
        # Question Paragraph
        qp = doc.add_paragraph()
        qp.paragraph_format.space_before = Pt(6)
        qp.paragraph_format.space_after = Pt(3)
        qp.paragraph_format.line_spacing = 1.15

        # Question Number & Text
        q_num = qp.add_run(f"{item['num']}. ")
        q_num.font.name = 'Nirmala UI'
        q_num.font.size = Pt(10.5)
        q_num.font.bold = True

        q_txt = qp.add_run(item['q'])
        q_txt.font.name = 'Nirmala UI'
        q_txt.font.size = Pt(10.5)
        q_txt.font.bold = True

        # Extract 4 options
        opts = item.get('options', [])
        opt_a = opts[0] if len(opts) > 0 else ""
        opt_b = opts[1] if len(opts) > 1 else ""
        opt_c = opts[2] if len(opts) > 2 else ""
        opt_d = opts[3] if len(opts) > 3 else ""

        # 2x2 Borderless Table for Side-by-Side Options
        # Row 1: A) ...    B) ...
        # Row 2: C) ...    D) ...
        table = doc.add_table(rows=2, cols=2)
        set_table_no_borders(table)

        row1 = table.rows[0]
        c_a, c_b = row1.cells[0], row1.cells[1]
        c_a.width = col_width
        c_b.width = col_width
        set_cell_margins(c_a, top=20, bottom=20, left=40, right=40)
        set_cell_margins(c_b, top=20, bottom=20, left=40, right=40)

        p_a = c_a.paragraphs[0]
        p_a.paragraph_format.space_before = Pt(0)
        p_a.paragraph_format.space_after = Pt(2)
        r_a = p_a.add_run(opt_a)
        r_a.font.name = 'Nirmala UI'
        r_a.font.size = Pt(10)

        p_b = c_b.paragraphs[0]
        p_b.paragraph_format.space_before = Pt(0)
        p_b.paragraph_format.space_after = Pt(2)
        r_b = p_b.add_run(opt_b)
        r_b.font.name = 'Nirmala UI'
        r_b.font.size = Pt(10)

        row2 = table.rows[1]
        c_c, c_d = row2.cells[0], row2.cells[1]
        c_c.width = col_width
        c_d.width = col_width
        set_cell_margins(c_c, top=20, bottom=20, left=40, right=40)
        set_cell_margins(c_d, top=20, bottom=20, left=40, right=40)

        p_c = c_c.paragraphs[0]
        p_c.paragraph_format.space_before = Pt(0)
        p_c.paragraph_format.space_after = Pt(4)
        r_c = p_c.add_run(opt_c)
        r_c.font.name = 'Nirmala UI'
        r_c.font.size = Pt(10)

        p_d = c_d.paragraphs[0]
        p_d.paragraph_format.space_before = Pt(0)
        p_d.paragraph_format.space_after = Pt(4)
        r_d = p_d.add_run(opt_d)
        r_d.font.name = 'Nirmala UI'
        r_d.font.size = Pt(10)

    # Save docx
    doc.save(docx_path)
    print(f"[OK] Saved DOCX: {docx_path}")

    # Convert to PDF via Word COM Automation
    print("Converting to PDF via Word COM...")
    word = win32com.client.DispatchEx("Word.Application")
    word.Visible = False
    word.DisplayAlerts = 0
    try:
        if os.path.exists(pdf_path):
            try:
                os.remove(pdf_path)
            except Exception:
                pass
        doc_obj = word.Documents.Open(os.path.abspath(docx_path))
        doc_obj.SaveAs(os.path.abspath(pdf_path), FileFormat=17) # 17 = wdFormatPDF
        doc_obj.Close(False)
        print(f"[OK] Saved PDF: {pdf_path} ({os.path.getsize(pdf_path):,} bytes)")
    finally:
        word.Quit()

def dispatch_to_telegram(pdf_path):
    TOKEN = "8738700204:AAGfOluaOWUUp5HgZTN0ZXiwxbZrlDEituk"
    CHAT_ID = "6233907249"

    # Pure professional Tamil caption as requested
    caption = (
        "📝 <b>பொருளாதாரம் மாதிரித் தேர்வு – 5 (வினாத்தாள்)</b>\n\n"
        "• <b>வடிவமைப்பு</b>: அசல் தேர்வு வடிவம் (பக்கவாட்டு விருப்பங்கள் A-B மேலே, C-D கீழே)\n"
        "• <b>மொத்த வினாக்கள்</b>: 75\n"
        "• <b>நேரம்</b>: 1 மணி 30 நிமிடம் | <b>மதிப்பெண்கள்</b>: 75\n"
        "• <b>குறிப்பு</b>: மாதிரித் தேர்வு எழுதும் நோக்கில் விடைகள் மற்றும் விளக்கங்கள் முழுமையாக நீக்கப்பட்டுள்ளன.\n\n"
        "வெற்றி பெற வாழ்த்துகள்!"
    )

    print("Sending PDF to Telegram...")
    with open(pdf_path, "rb") as f:
        r = requests.post(
            f"https://api.telegram.org/bot{TOKEN}/sendDocument",
            data={"chat_id": CHAT_ID, "caption": caption, "parse_mode": "HTML"},
            files={"document": ("Economics_Test_5_Question_Paper_SideBySide.pdf", f, "application/pdf")},
            timeout=60
        )
        print("[OK] Telegram status:", r.status_code, r.json().get("ok"))

if __name__ == "__main__":
    desktop = r"C:\Users\mukil\Desktop"
    docx_file = os.path.join(desktop, "Economics_Test_5_Question_Paper_SideBySide.docx")
    pdf_file = os.path.join(desktop, "Economics_Test_5_Question_Paper_SideBySide.pdf")

    build_exam_paper(docx_file, pdf_file)
    dispatch_to_telegram(pdf_file)
    print("\nCompleted successfully! NO Drive storage used.")
