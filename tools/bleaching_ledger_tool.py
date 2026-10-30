import os
import sys
import logging
from datetime import datetime
from typing import Dict, Any, Optional

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

logging.basicConfig(level=logging.INFO)
logger = logging.getLogger("bleaching_ledger")

BASE_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
LOCAL_XLSX_PATH = os.path.join(BASE_DIR, "storage", "business", "sgc_bleaching_pl_master.xlsx")
GOOGLE_SHEET_FILE_ID = "1w3NVcCQv-pVV6MOz4sNcjD3VmSHGndyG4u4PKnG5FfQ"
GOOGLE_SHEET_URL = f"https://docs.google.com/spreadsheets/d/{GOOGLE_SHEET_FILE_ID}/edit?usp=drivesdk"
TOKEN_PATH = os.path.join(BASE_DIR, "storage", "vault", "google_drive_token.json")


def ensure_ledger_exists() -> openpyxl.Workbook:
    """Ensures the local Excel master workbook exists with all formatted tabs."""
    os.makedirs(os.path.dirname(LOCAL_XLSX_PATH), exist_ok=True)
    if os.path.exists(LOCAL_XLSX_PATH):
        try:
            return openpyxl.load_workbook(LOCAL_XLSX_PATH)
        except Exception as e:
            logger.warning(f"Failed to load existing workbook ({e}), recreating...")

    wb = openpyxl.Workbook()
    header_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")

    # 1. Summary Sheet
    ws1 = wb.active
    ws1.title = "Daily_PL_Summary"
    headers_summary = [
        "Date", "Total KGs Bleached", "Gross Revenue (Rs)", 
        "Chemicals Cost (Rs)", "Power EB (Rs)", "Labour Wages (Rs)", 
        "Maintenance & Misc (Rs)", "Total Expenses (Rs)", 
        "Net Daily Profit (Rs)", "Profit Margin %", "Remarks"
    ]
    ws1.append(headers_summary)
    fill_summary = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
    for col in range(1, len(headers_summary) + 1):
        cell = ws1.cell(row=1, column=col)
        cell.font = header_font
        cell.fill = fill_summary
        cell.alignment = Alignment(horizontal="center", vertical="center")

    # 2. Production Sheet
    ws2 = wb.create_sheet(title="Production_Revenue")
    headers_prod = [
        "Date", "Batch No", "Customer Name", "Material / Count", 
        "Weight (KGs)", "Rate per KG (Rs)", "Total Value (Rs)", "DC No", "Payment Status"
    ]
    ws2.append(headers_prod)
    fill_prod = PatternFill(start_color="2E75B6", end_color="2E75B6", fill_type="solid")
    for col in range(1, len(headers_prod) + 1):
        cell = ws2.cell(row=1, column=col)
        cell.font = header_font
        cell.fill = fill_prod
        cell.alignment = Alignment(horizontal="center", vertical="center")

    # 3. Expenses Sheet
    ws3 = wb.create_sheet(title="Daily_Expenses")
    headers_exp = [
        "Date", "Category", "Item Description", "Vendor / Payee", 
        "Qty / Unit", "Amount (Rs)", "Payment Mode", "Bill / Ref"
    ]
    ws3.append(headers_exp)
    fill_exp = PatternFill(start_color="C00000", end_color="C00000", fill_type="solid")
    for col in range(1, len(headers_exp) + 1):
        cell = ws3.cell(row=1, column=col)
        cell.font = header_font
        cell.fill = fill_exp
        cell.alignment = Alignment(horizontal="center", vertical="center")

    wb.save(LOCAL_XLSX_PATH)
    logger.info(f"Initialized new ledger workbook at {LOCAL_XLSX_PATH}")
    return wb


def sync_to_google_drive() -> bool:
    """Syncs the local Excel ledger to the live Google Spreadsheet in Mukil's Master Vault."""
    if not os.path.exists(LOCAL_XLSX_PATH):
        return False
    try:
        from google.oauth2.credentials import Credentials
        from google.auth.transport.requests import Request
        from googleapiclient.discovery import build
        from googleapiclient.http import MediaFileUpload

        if not os.path.exists(TOKEN_PATH):
            logger.error(f"Google Drive token not found at {TOKEN_PATH}")
            return False

        creds = Credentials.from_authorized_user_file(TOKEN_PATH)
        if creds.expired and creds.refresh_token:
            creds.refresh(Request())

        drive = build("drive", "v3", credentials=creds)
        media = MediaFileUpload(
            LOCAL_XLSX_PATH,
            mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            resumable=True
        )
        drive.files().update(
            fileId=GOOGLE_SHEET_FILE_ID,
            media_body=media,
            fields="id, name, webViewLink"
        ).execute()
        logger.info(f"Live Google Sheet updated successfully: {GOOGLE_SHEET_URL}")
        return True
    except Exception as e:
        logger.error(f"Failed to sync ledger to Google Drive: {e}")
        return False


def _update_daily_summary(wb: openpyxl.Workbook, target_date: str):
    """Recomputes summary values for a given date from Production and Expense tabs."""
    ws_summary = wb["Daily_PL_Summary"]
    ws_prod = wb["Production_Revenue"]
    ws_exp = wb["Daily_Expenses"]

    # Calculate production totals
    total_kgs = 0.0
    total_rev = 0.0
    for row in ws_prod.iter_rows(min_row=2, values_only=True):
        if row and str(row[0]).strip() == target_date:
            try:
                total_kgs += float(row[4] or 0)
                total_rev += float(row[6] or 0)
            except Exception:
                pass

    # Calculate expense totals by category
    chem_cost = 0.0
    power_cost = 0.0
    labour_cost = 0.0
    misc_cost = 0.0

    for row in ws_exp.iter_rows(min_row=2, values_only=True):
        if row and str(row[0]).strip() == target_date:
            cat = str(row[1] or "").strip().lower()
            try:
                amt = float(row[5] or 0)
            except Exception:
                amt = 0.0

            if "chem" in cat or "h2o2" in cat or "caustic" in cat:
                chem_cost += amt
            elif "power" in cat or "eb" in cat or "electr" in cat:
                power_cost += amt
            elif "labour" in cat or "labor" in cat or "wage" in cat:
                labour_cost += amt
            else:
                misc_cost += amt

    total_exp = chem_cost + power_cost + labour_cost + misc_cost
    net_profit = total_rev - total_exp
    margin_pct = f"{(net_profit / total_rev * 100):.1f}%" if total_rev > 0 else "0.0%"

    # Look for existing row
    found_row = None
    for r_idx, row in enumerate(ws_summary.iter_rows(min_row=2, values_only=True), start=2):
        if row and str(row[0]).strip() == target_date:
            found_row = r_idx
            break

    summary_values = [
        target_date, round(total_kgs, 1), round(total_rev, 2),
        round(chem_cost, 2), round(power_cost, 2), round(labour_cost, 2),
        round(misc_cost, 2), round(total_exp, 2), round(net_profit, 2),
        margin_pct, "Active Bleaching Run"
    ]

    if found_row:
        for c_idx, val in enumerate(summary_values, start=1):
            ws_summary.cell(row=found_row, column=c_idx, value=val)
    else:
        ws_summary.append(summary_values)


def record_production(
    customer: str,
    material: str,
    kgs: float,
    rate_per_kg: float,
    dc_no: str = "AUTO-DC",
    date_str: Optional[str] = None
) -> Dict[str, Any]:
    """Records a production batch and recalculates daily P&L."""
    date_str = date_str or datetime.now().strftime("%Y-%m-%d")
    wb = ensure_ledger_exists()
    ws_prod = wb["Production_Revenue"]

    batch_no = f"BATCH-{datetime.now().strftime('%m%d')}-{len(list(ws_prod.rows))}"
    total_val = round(float(kgs) * float(rate_per_kg), 2)

    ws_prod.append([
        date_str, batch_no, customer, material, 
        float(kgs), float(rate_per_kg), total_val, dc_no, "PENDING"
    ])

    _update_daily_summary(wb, date_str)
    wb.save(LOCAL_XLSX_PATH)
    synced = sync_to_google_drive()

    return {
        "status": "success",
        "date": date_str,
        "batch_no": batch_no,
        "customer": customer,
        "kgs": kgs,
        "rate": rate_per_kg,
        "total_revenue": total_val,
        "drive_synced": synced,
        "sheet_url": GOOGLE_SHEET_URL
    }


def record_expense(
    category: str,
    item: str,
    vendor: str,
    amount: float,
    qty_unit: str = "1 Lot",
    payment_mode: str = "UPI",
    bill_ref: str = "VOUCHER",
    date_str: Optional[str] = None
) -> Dict[str, Any]:
    """Records an expense item (Chemicals, EB, Labour, Misc) and updates Daily P&L."""
    date_str = date_str or datetime.now().strftime("%Y-%m-%d")
    wb = ensure_ledger_exists()
    ws_exp = wb["Daily_Expenses"]

    ws_exp.append([
        date_str, category, item, vendor,
        qty_unit, float(amount), payment_mode, bill_ref
    ])

    _update_daily_summary(wb, date_str)
    wb.save(LOCAL_XLSX_PATH)
    synced = sync_to_google_drive()

    return {
        "status": "success",
        "date": date_str,
        "category": category,
        "item": item,
        "vendor": vendor,
        "amount": amount,
        "drive_synced": synced,
        "sheet_url": GOOGLE_SHEET_URL
    }


def get_daily_summary(date_str: Optional[str] = None) -> Dict[str, Any]:
    """Retrieves the summary metrics for a given date."""
    date_str = date_str or datetime.now().strftime("%Y-%m-%d")
    wb = ensure_ledger_exists()
    ws_summary = wb["Daily_PL_Summary"]

    for row in ws_summary.iter_rows(min_row=2, values_only=True):
        if row and str(row[0]).strip() == date_str:
            return {
                "date": row[0],
                "total_kgs": row[1],
                "gross_revenue": row[2],
                "chemicals_cost": row[3],
                "power_eb": row[4],
                "labour_wages": row[5],
                "maintenance_misc": row[6],
                "total_expenses": row[7],
                "net_profit": row[8],
                "profit_margin": row[9],
                "sheet_url": GOOGLE_SHEET_URL
            }

    return {
        "date": date_str,
        "total_kgs": 0,
        "gross_revenue": 0.0,
        "total_expenses": 0.0,
        "net_profit": 0.0,
        "profit_margin": "0.0%",
        "sheet_url": GOOGLE_SHEET_URL,
        "message": "No entries logged yet for this date."
    }


def pull_and_sync_ledger() -> Dict[str, Any]:
    """
    Downloads the live Google Sheet from Google Drive (capturing any edits Mukil made on phone/web),
    re-calculates the Daily_PL_Summary for all recorded dates, saves locally, and syncs back to Drive.
    """
    try:
        from google.oauth2.credentials import Credentials
        from google.auth.transport.requests import Request
        from googleapiclient.discovery import build

        if not os.path.exists(TOKEN_PATH):
            return {"status": "error", "message": "Google Drive token missing"}

        creds = Credentials.from_authorized_user_file(TOKEN_PATH)
        if creds.expired and creds.refresh_token:
            creds.refresh(Request())

        drive = build("drive", "v3", credentials=creds)
        req = drive.files().export_media(
            fileId=GOOGLE_SHEET_FILE_ID,
            mimeType="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        )
        content = req.execute()
        with open(LOCAL_XLSX_PATH, "wb") as f:
            f.write(content)

        wb = openpyxl.load_workbook(LOCAL_XLSX_PATH)
        all_dates = set()
        if "Production_Revenue" in wb.sheetnames:
            for r in wb["Production_Revenue"].iter_rows(min_row=2, values_only=True):
                if r and r[0]:
                    all_dates.add(str(r[0]).strip())
        if "Daily_Expenses" in wb.sheetnames:
            for r in wb["Daily_Expenses"].iter_rows(min_row=2, values_only=True):
                if r and r[0]:
                    all_dates.add(str(r[0]).strip())

        for d in sorted(all_dates):
            _update_daily_summary(wb, d)

        wb.save(LOCAL_XLSX_PATH)
        sync_to_google_drive()

        latest_date = max(all_dates) if all_dates else datetime.now().strftime("%Y-%m-%d")
        summary = get_daily_summary(latest_date)
        return {
            "status": "success",
            "message": f"Successfully pulled live edits from Google Sheet and recalculated P&L for {len(all_dates)} dates.",
            "latest_date": latest_date,
            "summary": summary,
            "sheet_url": GOOGLE_SHEET_URL
        }
    except Exception as e:
        logger.error(f"Failed to pull and sync ledger: {e}")
        return {"status": "error", "message": str(e)}


if __name__ == "__main__":
    print("Testing Bleaching Ledger Tool...")
    # Add test batch
    res_p = record_production("Sowbhagiya Home Textiles", "30s Grey Cotton Yarn", 2000, 12.0)
    print("Production logged:", res_p)
    # Add test expense
    res_e = record_expense("Chemicals", "H2O2 50% (2 Carboys)", "Lakshmi Chemicals", 4800.0)
    print("Expense logged:", res_e)
    # Get summary
    summary = get_daily_summary()
    print("Daily Summary:", summary)
